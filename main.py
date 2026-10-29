import undetected_chromedriver as uc
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.common.exceptions import TimeoutException, StaleElementReferenceException
import time
import openpyxl
import random
import pickle
import re
import subprocess
import os

CONVERSATION_URL = "https://www.facebook.com/messages/t/1840951729334287"

def get_random_text_from_excel(excel_file, column_index=2):
    workbook = openpyxl.load_workbook(excel_file)
    sheet = workbook.active
    num_rows = sheet.max_row
    random_row = random.randint(1, num_rows)
    random_text = sheet.cell(row=random_row, column=column_index).value
    workbook.close()
    return random_text

def get_chrome_version():
    for cmd in ['google-chrome --version', 'chromium-browser --version', 'chromium --version']:
        try:
            out = subprocess.check_output(cmd.split()).decode()
            match = re.search(r'(\d+)\.', out)
            if match:
                return int(match.group(1))
        except:
            pass
    return None

chrome_ver = get_chrome_version()
print(f"Detected Chrome version: {chrome_ver}")

random_text = get_random_text_from_excel('zeker.xlsx')
random_text = random_text.replace('_x000D_', ' ').replace('_x000A_', ' ').replace('\n', ' ').replace('\r', ' ')
print(f"Selected text (len={len(random_text)}): {random_text[:80]}...")

options = uc.ChromeOptions()
options.add_argument('--no-sandbox')
options.add_argument('--disable-dev-shm-usage')
options.add_argument('--window-size=1920,1080')
options.add_argument('--disable-gpu')
options.add_argument('--lang=ar')
if chrome_ver:
    driver = uc.Chrome(version_main=chrome_ver, use_subprocess=True, options=options)
else:
    driver = uc.Chrome(use_subprocess=True, options=options)
print("WebDriver started successfully.")

cookies = []
try:
    driver.get("https://www.facebook.com")
    time.sleep(3)

    with open('cookies.pkl', 'rb') as f:
        cookies = pickle.load(f)
    for cookie in cookies:
        try:
            driver.add_cookie(cookie)
        except Exception as e:
            print(f"Failed to add cookie {cookie.get('name')}: {e}")
    print("Cookies loaded successfully.")

    driver.get(CONVERSATION_URL)
    print(f"Navigated to the conversation URL.")

    time.sleep(12)

    current_url = driver.current_url
    print(f"Final URL: {current_url}")

    if "login" in current_url.lower():
        print("ERROR: Cookies are expired! Facebook redirected to login page.")
        print("Please update cookies.pkl with fresh cookies.")
        driver.save_screenshot("debug.png")
        with open("debug_page_source.html", "w") as f:
            f.write(driver.page_source)
        raise SystemExit(1)

    selectors = [
        (By.XPATH, "/html/body/div[1]/div/div[1]/div/div[3]/div/div/div[1]/div[1]/div[4]/div[2]/div/div/div/div[1]/div/div/div/div/div/div[2]/div/div/div/div[2]/div/div[2]/div/div[2]/div/div[1]/div[1]/p"),
        (By.CSS_SELECTOR, "div[contenteditable='true'][role='textbox']"),
        (By.CSS_SELECTOR, "p[contenteditable='true']"),
        (By.CSS_SELECTOR, "div[contenteditable='true']"),
    ]

    message_input = None
    for by, selector in selectors:
        try:
            message_input = WebDriverWait(driver, 5).until(
                EC.presence_of_element_located((by, selector))
            )
            print(f"Found element with selector: {selector[:60]}...")
            break
        except TimeoutException:
            continue

    if message_input is None:
        print("ERROR: Could not find message input with any selector!")
        driver.save_screenshot("debug.png")
        with open("debug_page_source.html", "w") as f:
            f.write(driver.page_source)
        print("Saved debug.png and debug_page_source.html")
        raise SystemExit(1)

    def find_input():
        for by, selector in selectors:
            try:
                return driver.find_element(by, selector)
            except:
                continue
        return None

    for attempt in range(3):
        try:
            message_input.click()
            time.sleep(0.5)
            message_input.send_keys(random_text)
            time.sleep(0.5)
            message_input.send_keys(Keys.ENTER)
            print("Sent the message.")
            break
        except StaleElementReferenceException:
            print(f"Stale element on attempt {attempt + 1}, re-locating...")
            time.sleep(2)
            message_input = find_input()
            if message_input is None:
                print("ERROR: Could not re-find message input after stale element!")
                driver.save_screenshot("debug.png")
                with open("debug_page_source.html", "w") as f:
                    f.write(driver.page_source)
                raise SystemExit(1)

    time.sleep(3)

except TimeoutException:
    print("Timed out waiting for the message input element.")
    driver.save_screenshot("debug.png")
    with open("debug_page_source.html", "w") as f:
        f.write(driver.page_source)
    print("Saved debug.png and debug_page_source.html")
    raise SystemExit(1)
except SystemExit:
    raise
except Exception as e:
    print(f"An error occurred: {e}")
    driver.save_screenshot("debug.png")
    raise SystemExit(1)
finally:
    try:
        new_cookies = driver.get_cookies()
        with open('cookies.pkl', 'wb') as f:
            pickle.dump(new_cookies, f)
        print(f"Cookies refreshed. Total: {len(new_cookies)}")
    except:
        pass
    driver.quit()
    print("WebDriver quit successfully.")
